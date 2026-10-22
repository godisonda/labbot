import logging
import re
import shutil
import subprocess
from pathlib import Path
from typing import Dict, Optional

logger = logging.getLogger("mtuci-labbot.compile")


def _apply_gost_word_styling(docx_path: Path):
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor, Cm
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls
    except ImportError:
        logger.warning("Библиотека python-docx не установлена.")
        return

    doc = Document(str(docx_path))
    black = RGBColor(0, 0, 0)

    # 1. Поля документа
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(1.5)

    on_title_page = True
    title_headings = 0

    # 2. Форматирование абзацев
    for p in doc.paragraphs:
        text = p.text.strip()
        style_name = p.style.name.lower() if p.style else ""
        is_h1 = "heading 1" in style_name

        # Первый H1 — «ОТЧЕТ ПО … РАБОТЕ» на титуле, второй — заголовок самого отчёта.
        if on_title_page and is_h1:
            title_headings += 1
            if title_headings >= 2:
                on_title_page = False
                p.paragraph_format.page_break_before = True

        if on_title_page:
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.left_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4)

            if any(k in text.lower() for k in ["выполнил", "проверил"]):
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.color.rgb = black
                if "отчет" in text.lower() or "«разработка" in text.lower():
                    run.font.size = Pt(14)
                    run.font.bold = True
                elif "московский технический" in text.lower():
                    run.font.size = Pt(12)
                    run.font.bold = True
                else:
                    run.font.size = Pt(11)
            continue

        is_h1 = "heading 1" in style_name
        is_h2 = "heading 2" in style_name
        is_h3 = "heading 3" in style_name
        is_numbered_header = bool(re.match(r"^\d+\.?\s+[А-Яа-я]", text) or re.match(r"^\d+\.\d+\s+[А-Яа-я]", text))
        is_header = is_h1 or is_h2 or is_h3 or is_numbered_header

        if is_header:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.left_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = Pt(12) if (is_h1 or not "." in text[:4]) else Pt(8)
            p.paragraph_format.space_after = Pt(4)
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.color.rgb = black
                run.font.bold = True
                run.font.size = Pt(14) if (is_h1 or not "." in text[:4]) else Pt(12)
            continue

        is_code = "source code" in style_name or "code" in style_name or (p.style and "verbatim" in p.style.name.lower())
        if is_code:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs:
                run.font.name = "Courier New"
                run.font.size = Pt(9.5)
                run.font.color.rgb = black
            continue

        if text.startswith("Рисунок") or text.startswith("Рис."):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(11)
                run.font.color.rgb = black
            continue

        if text.startswith("Таблица"):
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(11)
                run.font.color.rgb = black
            continue

        # Обычный текст
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(1.25)
        p.paragraph_format.left_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)

        for run in p.runs:
            run.font.name = "Times New Roman"
            run.font.color.rgb = black
            if run.font.size is None:
                run.font.size = Pt(12)

    # 3. Форматирование таблиц: границы, центрирование и шрифты
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        # Добавляем сетку границ к таблице Word
        try:
            tblPr = table._tbl.tblPr
            borders_xml = parse_xml(
                f'<w:tblBorders {nsdecls("w")}>\n'
                f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="A0A0A0"/>\n'
                f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="A0A0A0"/>\n'
                f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="D0D0D0"/>\n'
                f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="D0D0D0"/>\n'
                f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="A0A0A0"/>\n'
                f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="A0A0A0"/>\n'
                f'</w:tblBorders>'
            )
            tblPr.append(borders_xml)
        except Exception:
            pass

        for row in table.rows:
            for cell in row.cells:
                for cp in cell.paragraphs:
                    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cp.paragraph_format.first_line_indent = Cm(0)
                    cp.paragraph_format.left_indent = Cm(0)
                    cp.paragraph_format.space_before = Pt(2)
                    cp.paragraph_format.space_after = Pt(2)
                    cp.paragraph_format.line_spacing = 1.0
                    for run in cp.runs:
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(10.5)
                        run.font.color.rgb = black

    doc.save(str(docx_path))
    logger.info("Стили ГОСТ и рамки таблиц успешно применены к: %s", docx_path.name)


def compile_latex(build_dir: Path, timeout: int = 180) -> bool:
    build_dir = Path(build_dir).resolve()
    main_tex = build_dir / "main.tex"
    if not main_tex.exists():
        logger.error("Файл main.tex не найден в %s", build_dir)
        return False

    pdf_path = build_dir / "main.pdf"
    if pdf_path.exists():
        pdf_path.unlink()

    cmd = ["latexmk", "-xelatex", "-interaction=nonstopmode", "-file-line-error", "-outdir=.", "main.tex"]

    logger.info("Компиляция PDF через XeLaTeX в %s...", build_dir)
    try:
        proc = subprocess.run(cmd, cwd=str(build_dir), stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, timeout=timeout)
        log_file = build_dir / "latexmk.log"
        log_file.write_text(proc.stdout, encoding="utf-8")

        if pdf_path.exists() and pdf_path.stat().st_size > 1000:
            logger.info("PDF успешно собран: %s", pdf_path)
            return True

        xdv_path = build_dir / "main.xdv"
        if xdv_path.exists():
            subprocess.run(["xdvipdfmx", "-o", "main.pdf", "main.xdv"], cwd=str(build_dir), timeout=60)
            if pdf_path.exists() and pdf_path.stat().st_size > 1000:
                return True

        subprocess.run(["xelatex", "-interaction=nonstopmode", "main.tex"], cwd=str(build_dir), timeout=120)
        return pdf_path.exists() and pdf_path.stat().st_size > 1000

    except Exception as e:
        logger.error("Ошибка XeLaTeX: %s", e)
        return False


def compile_docx(build_dir: Path, lab_dir: Path, timeout: int = 60) -> bool:
    build_dir = Path(build_dir).resolve()
    word_md = build_dir / "report_for_word.md"
    if not word_md.exists():
        word_md = build_dir / "report_combined.md"

    if not word_md.exists():
        return False

    docx_path = build_dir / "report.docx"
    if docx_path.exists():
        docx_path.unlink()

    pandoc_cmd = shutil.which("pandoc")
    if not pandoc_cmd:
        logger.warning("Pandoc не установлен!")
        return False

    cmd = [
        pandoc_cmd,
        str(word_md.name),
        "-o",
        "report.docx",
        "--resource-path=.",
        "--from=markdown+pipe_tables",
        "--to=docx",
    ]

    logger.info("Сборка Word (.docx) через Pandoc в %s...", build_dir)
    try:
        proc = subprocess.run(cmd, cwd=str(build_dir), stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, timeout=timeout)
        if docx_path.exists() and docx_path.stat().st_size > 1000:
            _apply_gost_word_styling(docx_path)
            reports_dir = lab_dir / "reports"
            reports_dir.mkdir(parents=True, exist_ok=True)
            shutil.copy2(docx_path, reports_dir / "report.docx")
            logger.info("Итоговый Word готов: %s", reports_dir / "report.docx")
            return True
        else:
            logger.error("Ошибка Pandoc: %s", proc.stderr)
            return False
    except Exception as e:
        logger.error("Ошибка Pandoc: %s", e)
        return False


def compile_all(build_dir: Path, lab_dir: Path, client: Optional[object] = None) -> Dict[str, bool]:
    results = {}
    results["pdf"] = compile_latex(build_dir)
    if results["pdf"]:
        pdf_build = build_dir / "main.pdf"
        reports_dir = lab_dir / "reports"
        reports_dir.mkdir(parents=True, exist_ok=True)
        shutil.copy2(pdf_build, reports_dir / "main.pdf")

    results["docx"] = compile_docx(build_dir, lab_dir)
    return results
